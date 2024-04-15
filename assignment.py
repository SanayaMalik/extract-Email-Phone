import os
import re
import PyPDF2
import docx
import xlwt  # Import xlwt library for writing Excel files in .xls format

def extract_text_from_pdf(pdf_path):
    text = ""
    with open(pdf_path, 'rb') as f:
        reader = PyPDF2.PdfReader(f)
        for page_num in range(len(reader.pages)):
            text += reader.pages[page_num].extract_text()
    return text

def extract_text_from_docx(docx_path):
    doc = docx.Document(docx_path)
    text = ""
    for para in doc.paragraphs:
        text += para.text
    return text

def extract_text_from_csv(csv_path):
    with open(csv_path, 'r') as file:
        lines = file.readlines()
        text = ' '.join(lines)
    return text

def extract_emails_and_phone_numbers(text):
    emails = re.findall(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b', text)
    phones = re.findall(r'(?:(?:\+?([1-9]|[0-9][0-9]|[0-9][0-9][0-9])\s*(?:[.-]\s*)?)?(?:\([2-9]0[1-9]\)|[2-9]0[1-9])\s*(?:[.-]\s*)?)?([2-9]0[1-9])\s*(?:[.-]\s*)?([0-9]{4})(?:\s*(?:#|x\.?|ext\.?|extension)\s*(\d+))?', text)
    emails = '\n'.join(emails)
    phones = [''.join(filter(str.isdigit, ''.join(phone))) for phone in phones if ''.join(filter(str.isdigit, ''.join(phone)))]
    phones = '\n'.join(phones)
    return emails, phones

def process_file(file_path):
    if file_path.endswith('.pdf'):
        text = extract_text_from_pdf(file_path)
    elif file_path.endswith('.docx'):
        text = extract_text_from_docx(file_path)
    elif file_path.endswith('.csv'):
        text = extract_text_from_csv(file_path)
    else:
        return "", ""  # Skip non-PDF, non-DOCX, and non-CSV files
    return extract_emails_and_phone_numbers(text)

def process_folder(folder_path):
    emails_and_phones = []
    if os.path.exists(folder_path):
        for filename in os.listdir(folder_path):
            file_path = os.path.join(folder_path, filename)
            file_emails, file_phones = process_file(file_path)
            emails_and_phones.append((file_path, file_emails, file_phones))
    else:
        print("Folder path does not exist.")
    return emails_and_phones

def write_to_excel(data, excel_path):
    workbook = xlwt.Workbook()  # Create a workbook
    worksheet = workbook.add_sheet('Sheet1')  # Add a worksheet
    worksheet.write(0, 0, "File")  # Write headers
    worksheet.write(0, 1, "Email")
    worksheet.write(0, 2, "Phone Number")
    row = 1
    for file_path, emails, phones in data:
        worksheet.write(row, 0, file_path)  # Write data
        worksheet.write(row, 1, emails)
        worksheet.write(row, 2, phones)
        row += 1
    workbook.save(excel_path)  # Save workbook to Excel file
